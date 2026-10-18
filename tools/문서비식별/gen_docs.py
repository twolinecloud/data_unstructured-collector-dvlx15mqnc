# 문서 비식별 시험 문서 생성기 (1단계 탐침용 — 포맷별 1쪽)
#
# 정답을 아는 문서만 채점한다. 그래서 PII 는 여기 한 곳에 박아 두고
# 모든 포맷이 같은 값을 쓴다 — 포맷 간 차이가 '문서 탓' 인지 '값 탓' 인지 섞이면 못 읽는다.
import io, json, os, sys

OUT = sys.argv[1] if len(sys.argv) > 1 else "."
os.makedirs(OUT, exist_ok=True)

def rrn_checksum(d12):
    """주민등록번호 검증숫자. AI-R 의 주민 패턴(6a78df5a…)은 **체크섬을 검사한다** —
    2026-10-18 실측. 체크섬 안 맞는 가상 번호를 쓰면 영원히 미탐으로 보이고,
    그걸 모르면 'AI-R 이 주민번호를 못 잡는다' 는 틀린 결론이 나간다(한 번 그랬다)."""
    w = [2, 3, 4, 5, 6, 7, 8, 9, 2, 3, 4, 5]
    return str((11 - sum(int(c) * k for c, k in zip(d12, w)) % 11) % 10)


_RRN_FRONT, _RRN_MID = "900101", "123456"
_RRN_OK = _RRN_FRONT + _RRN_MID + rrn_checksum(_RRN_FRONT + _RRN_MID)

# 가상 PII — 교정번호는 DMY 대역, 실데이터 금지
PII = {
    "name1": "홍길동",
    "name2": "김영희",
    "name3": "박철수",
    "phone": "010-1234-5678",
    "phone_land": "02-3456-7890",
    "rrn": _RRN_OK[:6] + "-" + _RRN_OK[6:],            # 체크섬 맞음 — 잡혀야 한다
    "rrn_bad": "900101-1234567",                     # 체크섬 틀림 — 함정 기록용(안 잡힌다)
    "rrn_nohyphen": _RRN_OK,                          # 하이픈 없음 — 체크섬 맞아도 안 잡힌다(진짜 구멍)
    "acct": "110-234-567890",
    "addr": "서울특별시 강남구 테헤란로 123",
    "org": "서울남부교정시설",
    "email": "hong.gildong@example.com",
}

# 자리별로 심는다 — 추출 누락은 자리마다 다르게 난다(10-08: hwpx 가 '표 뒤 문단' 을 잃었다)
PLACES = [
    ("본문", "수용자 {name1} (전화 {phone}) 은 {org} 에 수용 중이며 주소는 {addr} 이다."),
    ("본문2", "보호자 {name2} 의 연락처는 {phone_land} 이고 주민등록번호는 {rrn} 이다."),
    ("본문3", "환불 계좌는 {acct} 이며 담당자 이메일은 {email} 이다."),
    ("본문4-주민변형", "대조용 — 체크섬 틀린 번호 {rrn_bad} · 하이픈 없는 번호 {rrn_nohyphen}."),
]

ANSWERS = []


def ans(place, kind, value, note=""):
    ANSWERS.append({"place": place, "kind": kind, "value": value, "note": note})


def body_lines():
    out = []
    for place, tpl in PLACES:
        line = tpl.format(**PII)
        out.append((place, line))
    return out


def register_body():
    for place, tpl in PLACES:
        for k, v in PII.items():
            if "{" + k + "}" in tpl:
                ans(place, k, v)


# ---------------------------------------------------------------- docx
def make_docx(path):
    import docx
    from docx.shared import Pt

    d = docx.Document()
    d.add_heading("문서 비식별 시험 문서 (docx)", level=1)
    for place, line in body_lines():
        d.add_paragraph(line)

    # 표 ①  단순 2열 — 이게 되는지가 기준선
    d.add_paragraph("표1. 단순 2열")
    t = d.add_table(rows=3, cols=2)
    t.style = "Table Grid"
    rows = [("성명", PII["name1"]), ("연락처", PII["phone"]), ("주소", PII["addr"])]
    for i, (k, v) in enumerate(rows):
        t.cell(i, 0).text = k
        t.cell(i, 1).text = v
    ans("표-단순", "name1", PII["name1"])
    ans("표-단순", "phone", PII["phone"])
    ans("표-단순", "addr", PII["addr"])

    # 표 ②  값이 세 칸에 나뉨 — 분절되면 전화 패턴이 안 맞는다
    d.add_paragraph("표2. 값 분할(지역번호/국번/번호)")
    t2 = d.add_table(rows=2, cols=3)
    t2.style = "Table Grid"
    for j, h in enumerate(["지역", "국번", "번호"]):
        t2.cell(0, j).text = h
    for j, v in enumerate(["010", "1234", "5678"]):
        t2.cell(1, j).text = v
    ans("표-값분할", "phone_split", "010 / 1234 / 5678", "셀 3칸으로 쪼갠 010-1234-5678")

    # 표 ③  셀 안 줄바꿈 — 한 값이 두 줄로
    d.add_paragraph("표3. 셀 안 줄바꿈")
    t3 = d.add_table(rows=1, cols=2)
    t3.style = "Table Grid"
    t3.cell(0, 0).text = "연락처"
    c = t3.cell(0, 1).paragraphs[0]
    c.add_run("010-")
    c.add_run().add_break()
    c.add_run("9876-5432")
    ans("표-셀내줄바꿈", "phone_wrapped", "010-9876-5432", "셀 안에서 두 줄로 끊은 번호")

    # 표 ④  병합 셀 + 이름만 있는 칸(10-08 에서 가장 많이 놓친 자리)
    d.add_paragraph("표4. 병합 + 이름만")
    t4 = d.add_table(rows=2, cols=3)
    t4.style = "Table Grid"
    t4.cell(0, 0).merge(t4.cell(0, 1))
    t4.cell(0, 0).text = "면회 신청자"
    t4.cell(0, 2).text = "관계"
    t4.cell(1, 0).text = PII["name2"]
    t4.cell(1, 1).text = PII["name3"]
    t4.cell(1, 2).text = "지인"
    ans("표-이름만", "name2", PII["name2"], "문장 없이 이름만 있는 셀")
    ans("표-이름만", "name3", PII["name3"], "문장 없이 이름만 있는 셀")

    # 표 바로 뒤 문단 — hwpx 기지 결함 재현 자리
    d.add_paragraph("표 바로 뒤 문단: 담당 교도관 {name1}, 연락처 {phone_land}".format(**PII))
    ans("표뒤문단", "name1", PII["name1"], "hwpx 가 잃는 자리")
    ans("표뒤문단", "phone_land", PII["phone_land"], "hwpx 가 잃는 자리")

    # 머리말·꼬리말
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "머리말 — 작성자 {name1}".format(**PII)
    sec.footer.paragraphs[0].text = "꼬리말 — 문의 {phone}".format(**PII)
    ans("머리말", "name1", PII["name1"])
    ans("꼬리말", "phone", PII["phone"])

    # 문서 속성(10-08: docProps 도 검출 대상)
    d.core_properties.author = PII["name1"]
    d.core_properties.comments = "연락처 " + PII["phone"]
    ans("문서속성", "name1", PII["name1"])
    ans("문서속성", "phone", PII["phone"])

    d.save(path)


# ---------------------------------------------------------------- xlsx
def make_xlsx(path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "본문"
    r = 1
    for place, line in body_lines():
        ws.cell(row=r, column=1, value=line)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="성명"); ws.cell(row=r, column=2, value=PII["name1"]); r += 1
    ws.cell(row=r, column=1, value="연락처"); ws.cell(row=r, column=2, value=PII["phone"]); r += 1
    ws.cell(row=r, column=1, value="주소"); ws.cell(row=r, column=2, value=PII["addr"]); r += 2
    # 값 분할
    for j, v in enumerate(["010", "1234", "5678"]):
        ws.cell(row=r, column=1 + j, value=v)
    r += 2
    # 병합 + 이름만
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    ws.cell(row=r, column=1, value="면회 신청자"); r += 1
    ws.cell(row=r, column=1, value=PII["name2"]); ws.cell(row=r, column=2, value=PII["name3"]); r += 2
    ws.cell(row=r, column=1, value="표 뒤 문단: 담당 교도관 {name1}, 연락처 {phone_land}".format(**PII))
    ws2 = wb.create_sheet("둘째시트")
    ws2["A1"] = "둘째 시트 — 수용자 {name3} 주민 {rrn}".format(**PII)
    wb.properties.creator = PII["name1"]
    wb.save(path)


# ---------------------------------------------------------------- pptx
def make_pptx(path):
    from pptx import Presentation
    from pptx.util import Inches, Pt

    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "문서 비식별 시험 문서 (pptx)"
    tf = s.placeholders[1].text_frame
    for place, line in body_lines():
        tf.add_paragraph().text = line

    # 표 슬라이드
    s2 = p.slides.add_slide(p.slide_layouts[5])
    s2.shapes.title.text = "표 — 단순 / 값분할 / 이름만"
    t = s2.shapes.add_table(4, 3, Inches(0.5), Inches(1.6), Inches(9), Inches(2.4)).table
    hdr = ["구분", "값", "비고"]
    for j, h in enumerate(hdr):
        t.cell(0, j).text = h
    data = [
        ("성명", PII["name1"], ""),
        ("연락처", PII["phone"], ""),
        ("면회", PII["name2"], PII["name3"]),
    ]
    for i, row in enumerate(data, start=1):
        for j, v in enumerate(row):
            t.cell(i, j).text = v

    # 도형 글상자
    s3 = p.slides.add_slide(p.slide_layouts[6])
    tb = s3.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(9), Inches(1.0))
    tb.text_frame.text = "도형 글상자 — 주소 {addr} / 계좌 {acct}".format(**PII)

    # 발표자 노트
    s.notes_slide.notes_text_frame.text = "노트 — 담당 {name3} {phone_land}".format(**PII)
    p.core_properties.author = PII["name1"]
    p.save(path)


if __name__ == "__main__":
    register_body()
    make_docx(os.path.join(OUT, "deid_probe.docx"))
    make_xlsx(os.path.join(OUT, "deid_probe.xlsx"))
    make_pptx(os.path.join(OUT, "deid_probe.pptx"))
    with io.open(os.path.join(OUT, "answers.json"), "w", encoding="utf-8") as f:
        json.dump({"pii": PII, "answers": ANSWERS}, f, ensure_ascii=False, indent=2)
    for n in ("deid_probe.docx", "deid_probe.xlsx", "deid_probe.pptx", "answers.json"):
        p = os.path.join(OUT, n)
        print(n, os.path.getsize(p), "bytes")
    print("정답 항목", len(ANSWERS), "건 (docx 기준 자리별)")
